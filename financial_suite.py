#!/usr/bin/env python3
"""
Financial Reporting & Document Processing Suite

Reports, Data Extraction, Reconciliation  

Install
-------
    pip install pandas openpyxl pdfplumber pytesseract pillow

Run the UI
----------
    python financial_suite.py


"""

from __future__ import annotations

import argparse
import difflib
import re
import sys
import threading
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
except ImportError:  # headless environments
    tk = None

# ----------------------------------------------------------------------------
# Shared helpers
# ----------------------------------------------------------------------------
FONT = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1F3864")
SECTION_FILL = PatternFill("solid", start_color="D9E1F2")
TOTAL_BORDER = Border(top=Side(style="thin"), bottom=Side(style="double"))
MONEY_FMT = '$#,##0.00;($#,##0.00);"-"'


def clean_amount(value) -> float:
    """'$1,234.50', '(200.00)', '-45' -> float. Blank/invalid -> NaN."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return float("nan")
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip()
    if not s:
        return float("nan")
    neg = s.startswith("(") and s.endswith(")") or s.startswith("-")
    s = re.sub(r"[^\d.]", "", s)
    if not s:
        return float("nan")
    try:
        n = float(s)
    except ValueError:
        return float("nan")
    return -n if neg else n


def read_table(path: str) -> pd.DataFrame:
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xlsm", ".xls"):
        return pd.read_excel(p)
    return pd.read_csv(p)


def load_transactions(path: str) -> pd.DataFrame:
    """Load a ledger/bank file and normalise to Date, Description, Category,
    Amount, Reference. Accepts either an Amount column (income +, expense -)
    or separate Debit/Credit columns."""
    df = read_table(path)
    df.columns = [str(c).strip().lower() for c in df.columns]
    aliases = {
        "date": ["date", "transaction date", "posted date", "posting date", "txn date"],
        "description": ["description", "memo", "details", "narrative", "payee", "name"],
        "category": ["category", "account", "gl account", "type"],
        "reference": ["reference", "ref", "check number", "check #", "transaction id", "id"],
        "amount": ["amount", "amt", "value"],
        "debit": ["debit", "withdrawal", "withdrawals"],
        "credit": ["credit", "deposit", "deposits"],
    }
    out = pd.DataFrame()
    for target, names in aliases.items():
        for n in names:
            if n in df.columns:
                out[target] = df[n]
                break
    if "date" not in out:
        raise ValueError("No date column found (expected e.g. 'Date').")
    if "amount" in out:
        out["Amount"] = out["amount"].map(clean_amount)
    elif "debit" in out or "credit" in out:
        d = out["debit"].map(clean_amount).fillna(0) if "debit" in out else 0
        c = out["credit"].map(clean_amount).fillna(0) if "credit" in out else 0
        out["Amount"] = c - d
    else:
        raise ValueError("No Amount (or Debit/Credit) column found.")
    res = pd.DataFrame({
        "Date": pd.to_datetime(out["date"], errors="coerce"),
        "Description": out.get("description", "").astype(str) if "description" in out else "",
        "Category": out["category"].astype(str) if "category" in out else "Uncategorized",
        "Reference": out["reference"].astype(str).str.strip() if "reference" in out else "",
        "Amount": out["Amount"],
    })
    res = res.dropna(subset=["Date", "Amount"]).reset_index(drop=True)
    res["Category"] = res["Category"].replace({"nan": "Uncategorized", "": "Uncategorized"})
    res["Reference"] = res["Reference"].replace({"nan": ""})
    return res


def style_header(ws, row: int, ncols: int):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def autofit(ws, min_w=10, max_w=60):
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, min_w), max_w)


def write_df(ws, df: pd.DataFrame, start_row=1, money_cols=(), date_cols=()):
    for j, col in enumerate(df.columns, 1):
        ws.cell(row=start_row, column=j, value=col)
    style_header(ws, start_row, len(df.columns))
    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, pd.Timestamp):
                v = v.to_pydatetime()
            elif pd.isna(v) if not isinstance(v, (list, dict)) else False:
                v = None
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = Font(name=FONT)
            name = df.columns[j - 1]
            if name in money_cols:
                cell.number_format = MONEY_FMT
            if name in date_cols:
                cell.number_format = "yyyy-mm-dd"
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    autofit(ws)


# ----------------------------------------------------------------------------
# 1. Financial statement automation
# ----------------------------------------------------------------------------
def add_period(df: pd.DataFrame, freq: str) -> pd.DataFrame:
    """Adds PeriodStart (sortable) and Period (label) columns."""
    df = df.copy()
    freq = freq.lower()
    if freq == "weekly":
        start = df["Date"].dt.to_period("W-SUN").dt.start_time
        label = "Wk of " + start.dt.strftime("%Y-%m-%d")
    elif freq == "monthly":
        start = df["Date"].dt.to_period("M").dt.start_time
        label = start.dt.strftime("%b %Y")
    elif freq == "quarterly":
        q = df["Date"].dt.to_period("Q")
        start = q.dt.start_time
        label = q.dt.year.astype(str) + " Q" + q.dt.quarter.astype(str)
    else:
        raise ValueError(f"Unknown frequency: {freq}")
    df["PeriodStart"], df["Period"] = start, label
    return df


def build_financial_statement(df: pd.DataFrame, freq: str, out_path: str,
                              company: str = "My Company") -> pd.DataFrame:
    """Write an Income Statement workbook (Excel formulas for every total, plus
    a chart and the source transactions). Returns a Period/Revenue/Expenses/Net
    summary DataFrame for on-screen preview."""
    df = add_period(df, freq)
    order = (df[["PeriodStart", "Period"]].drop_duplicates().sort_values("PeriodStart")["Period"].tolist())
    rev = df[df.Amount > 0].pivot_table(index="Category", columns="Period", values="Amount",
                                        aggfunc="sum", fill_value=0).reindex(columns=order, fill_value=0)
    exp = df[df.Amount < 0].assign(Amount=lambda d: d.Amount.abs()).pivot_table(
        index="Category", columns="Period", values="Amount", aggfunc="sum", fill_value=0
    ).reindex(columns=order, fill_value=0)

    wb = Workbook()
    ws = wb.active
    ws.title = "Income Statement"
    ncols = len(order) + 2  # label + periods + total
    ws["A1"] = company
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = f"Income Statement - {freq.capitalize()}"
    ws["A2"].font = Font(name=FONT, bold=True, size=11)
    ws["A3"] = (f"Period covered: {df.Date.min():%Y-%m-%d} to {df.Date.max():%Y-%m-%d}   |   "
                f"Generated {datetime.now():%Y-%m-%d %H:%M}")
    ws["A3"].font = Font(name=FONT, italic=True, size=9)

    hdr = 5
    ws.cell(row=hdr, column=1, value="Line item")
    for j, p in enumerate(order, 2):
        ws.cell(row=hdr, column=j, value=p)
    ws.cell(row=hdr, column=ncols, value="Total")
    style_header(ws, hdr, ncols)

    def section(title, data, row):
        ws.cell(row=row, column=1, value=title).font = Font(name=FONT, bold=True)
        for c in range(1, ncols + 1):
            ws.cell(row=row, column=c).fill = SECTION_FILL
        row += 1
        first = row
        for cat, vals in data.iterrows():
            ws.cell(row=row, column=1, value=f"   {cat}").font = Font(name=FONT)
            for j, p in enumerate(order, 2):
                c = ws.cell(row=row, column=j, value=round(float(vals[p]), 2))
                c.number_format, c.font = MONEY_FMT, Font(name=FONT, color="0000FF")  # blue = input data
            t = ws.cell(row=row, column=ncols, value=f"=SUM(B{row}:{get_column_letter(ncols - 1)}{row})")
            t.number_format, t.font = MONEY_FMT, Font(name=FONT)
            row += 1
        last = row - 1
        ws.cell(row=row, column=1, value=f"Total {title}").font = Font(name=FONT, bold=True)
        for j in range(2, ncols + 1):
            L = get_column_letter(j)
            f = f"=SUM({L}{first}:{L}{last})" if last >= first else 0
            c = ws.cell(row=row, column=j, value=f)
            c.number_format, c.font, c.border = MONEY_FMT, Font(name=FONT, bold=True), TOTAL_BORDER
        return row

    rev_tot = section("Revenue", rev, hdr + 1)
    exp_tot = section("Expenses", exp, rev_tot + 2)
    net = exp_tot + 2
    ws.cell(row=net, column=1, value="Net Income").font = Font(name=FONT, bold=True)
    mrg = net + 1
    ws.cell(row=mrg, column=1, value="Net margin").font = Font(name=FONT, italic=True)
    for j in range(2, ncols + 1):
        L = get_column_letter(j)
        c = ws.cell(row=net, column=j, value=f"={L}{rev_tot}-{L}{exp_tot}")
        c.number_format, c.font, c.border = MONEY_FMT, Font(name=FONT, bold=True), TOTAL_BORDER
        m = ws.cell(row=mrg, column=j, value=f"=IF({L}{rev_tot}=0,0,{L}{net}/{L}{rev_tot})")
        m.number_format, m.font = "0.0%;(0.0%);-", Font(name=FONT, italic=True)

    # chart block (contiguous rows linking to totals)
    cb = mrg + 3
    ws.cell(row=cb - 1, column=1, value="Chart data").font = Font(name=FONT, bold=True, size=9)
    for k, (lbl, src) in enumerate([("Revenue", rev_tot), ("Expenses", exp_tot), ("Net Income", net)]):
        ws.cell(row=cb + k, column=1, value=lbl).font = Font(name=FONT, size=9)
        for j in range(2, ncols):
            c = ws.cell(row=cb + k, column=j, value=f"={get_column_letter(j)}{src}")
            c.number_format, c.font = MONEY_FMT, Font(name=FONT, size=9)
    ch = BarChart()
    ch.type, ch.title, ch.height, ch.width = "col", f"{freq.capitalize()} Revenue vs Expenses", 8, 18
    ch.add_data(Reference(ws, min_col=1, max_col=ncols - 1, min_row=cb, max_row=cb + 2), from_rows=True, titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, max_col=ncols - 1, min_row=hdr))
    ws.add_chart(ch, f"A{cb + 5}")

    note = cb + 24
    ws.cell(row=note, column=1,
            value="Notes: blue figures are aggregated from the 'Transactions' sheet; black cells are formulas. "
                  "Positive amounts = revenue, negative = expenses.").font = Font(name=FONT, italic=True, size=8)
    ws.column_dimensions["A"].width = 34
    for j in range(2, ncols + 1):
        ws.column_dimensions[get_column_letter(j)].width = 16
    ws.freeze_panes = ws.cell(row=hdr + 1, column=2)
    ws.sheet_view.showGridLines = False

    tx = wb.create_sheet("Transactions")
    write_df(tx, df[["Date", "Description", "Category", "Reference", "Amount", "Period"]],
             money_cols=("Amount",), date_cols=("Date",))

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)

    g = df.groupby(["PeriodStart", "Period"]).Amount
    summary = pd.DataFrame({
        "Period": [k[1] for k in g.sum().index],
        "Revenue": g.apply(lambda s: s[s > 0].sum()).values,
        "Expenses": g.apply(lambda s: -s[s < 0].sum()).values,
        "Net Income": g.sum().values,
    }).round(2)
    return summary


def generate_reports(ledger_path: str, freq: str, out_dir: str, company: str):
    df = load_transactions(ledger_path)
    freqs = ["weekly", "monthly", "quarterly"] if freq.lower() == "all" else [freq.lower()]
    stamp = datetime.now().strftime("%Y%m%d")
    results = {}
    for f in freqs:
        path = str(Path(out_dir) / f"income_statement_{f}_{stamp}.xlsx")
        results[f] = (path, build_financial_statement(df, f, path, company))
    return results


# ----------------------------------------------------------------------------
# 2. Data extraction (invoices, scanned checks)
# ----------------------------------------------------------------------------
MONEY = r"\$?\s*((?:\d{1,3}(?:,\d{3})+|\d+)\.\d{2})"
DATE_PAT = (r"(\d{4}-\d{2}-\d{2}|\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|"
            r"(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*\.?\s+\d{1,2},?\s+\d{4}|"
            r"\d{1,2}\s+(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*\.?,?\s+\d{4})")


def ocr_available() -> bool:
    try:
        import pytesseract
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def _ocr_image(img) -> str:
    import pytesseract
    from PIL import ImageOps
    img = ImageOps.autocontrast(ImageOps.grayscale(img))
    return pytesseract.image_to_string(img, config="--psm 6")


def extract_text(path: str) -> tuple[str, str]:
    """Return (text, method). Uses the PDF text layer when present, else OCR."""
    p = Path(path)
    if p.suffix.lower() == ".pdf":
        import pdfplumber
        text, pages = "", []
        with pdfplumber.open(p) as pdf:
            pages = list(pdf.pages)
            text = "\n".join((pg.extract_text() or "") for pg in pages)
            if len(text.strip()) >= 30:
                return text, "pdf-text"
            if not ocr_available():
                raise RuntimeError("Scanned PDF detected but Tesseract OCR is not installed.")
            ocr = [_ocr_image(pg.to_image(resolution=300).original) for pg in pages]
        return "\n".join(ocr), "ocr"
    from PIL import Image
    if not ocr_available():
        raise RuntimeError("Tesseract OCR is not installed (needed for images).")
    return _ocr_image(Image.open(p)), "ocr"


def _parse_date(s: str | None):
    if not s:
        return None
    try:
        from dateutil import parser
        return parser.parse(s, dayfirst=False).strftime("%Y-%m-%d")
    except Exception:
        return s


def _date_after(label: str, text: str):
    m = re.search(label + r"[^\n\d]{0,20}" + DATE_PAT, text, re.I)
    return _parse_date(m.group(1)) if m else None


def _amount_on_line(text: str, label_regex: str, exclude: str | None = None, last=True):
    hits = []
    for line in text.splitlines():
        if re.search(label_regex, line, re.I) and not (exclude and re.search(exclude, line, re.I)):
            nums = re.findall(MONEY, line)
            if nums:
                hits.append(clean_amount(nums[-1]))
    if not hits:
        return None
    return hits[-1] if last else hits[0]


def parse_invoice(text: str) -> dict:
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    vendor = next((l for l in lines if not re.search(r"invoice|bill to|^\d+$|^page", l, re.I) and len(l) > 2), None)
    inv_no = None
    for pat in (r"(?:invoice|inv)\s*(?:no\.?|number|num|#)\s*[:#.\-]?\s*([A-Z0-9][A-Z0-9\-/]{2,})",
                r"invoice\s*[:#]\s*([A-Z0-9][A-Z0-9\-/]{2,})"):
        m = re.search(pat, text, re.I)
        if m:
            inv_no = m.group(1)
            break
    total = (_amount_on_line(text, r"grand total|total due|amount due|balance due")
             or _amount_on_line(text, r"\btotal\b", exclude=r"sub\s*-?\s*total"))
    po = re.search(r"(?:p\.?o\.?|purchase order)\s*(?:no\.?|number|#)?\s*[:#]?\s*([A-Z0-9\-]{3,})", text, re.I)
    return {
        "Type": "Invoice",
        "Vendor": vendor,
        "Invoice #": inv_no,
        "Invoice Date": _date_after(r"(?:invoice\s*)?date", text),
        "Due Date": _date_after(r"due(?:\s*date)?", text),
        "PO #": po.group(1) if po else None,
        "Subtotal": _amount_on_line(text, r"sub\s*-?\s*total"),
        "Tax": _amount_on_line(text, r"\b(?:tax|vat|gst)\b"),
        "Total": total,
    }


def aba_valid(routing: str) -> bool:
    if not re.fullmatch(r"\d{9}", routing):
        return False
    d = list(map(int, routing))
    return (3 * (d[0] + d[3] + d[6]) + 7 * (d[1] + d[4] + d[7]) + (d[2] + d[5] + d[8])) % 10 == 0


def parse_check(text: str) -> dict:
    bank = next((l.strip() for l in text.splitlines() if re.search(r"\bbank\b", l, re.I)), None)
    payee = None
    m = re.search(r"pay\s*to\s*(?:the\s*)?(?:order\s*of)?\s*[:\-]?\s*([^\n$]{3,60})", text, re.I)
    if m:
        payee = re.sub(r"[_\.]{2,}.*", "", m.group(1)).strip(" _.-")
    amt = None
    m = re.search(r"\$\s*" + MONEY.replace(r"\$?\s*", "", 1), text)
    if m:
        amt = clean_amount(m.group(1))
    else:
        nums = [clean_amount(x) for x in re.findall(MONEY, text)]
        amt = max(nums) if nums else None
    routing = account = chk_no = None
    mm = re.search(r"(\d{9})\D{1,4}(\d{6,17})\D{1,4}(\d{3,6})", text)
    if mm:
        routing, account, chk_no = mm.groups()
    if not chk_no:
        m = re.search(r"check\s*(?:no\.?|number|#)\s*[:#]?\s*(\d{3,8})", text, re.I)
        chk_no = m.group(1) if m else None
    memo = re.search(r"memo\s*[:\-]?\s*([^\n]{2,60})", text, re.I)
    return {
        "Type": "Check",
        "Bank": bank,
        "Check #": chk_no,
        "Date": _date_after(r"date", text),
        "Pay To": payee,
        "Amount": amt,
        "Memo": memo.group(1).strip(" _.") if memo else None,
        "Routing #": routing,
        "Routing Valid (ABA)": aba_valid(routing) if routing else None,
        "Account #": account,
    }


def detect_type(text: str) -> str:
    if re.search(r"pay\s*to\s*the\s*order|\bpay\s*to\b", text, re.I) and not re.search(r"invoice", text, re.I):
        return "check"
    return "invoice"


def process_documents(paths: list[str], doc_type: str = "auto") -> tuple[pd.DataFrame, pd.DataFrame]:
    """Returns (invoices_df, checks_df). Each row also carries Source/Method/Error."""
    inv, chk = [], []
    for p in paths:
        base = {"Source": Path(p).name}
        try:
            text, method = extract_text(p)
            kind = doc_type if doc_type != "auto" else detect_type(text)
            rec = parse_check(text) if kind == "check" else parse_invoice(text)
            rec.update(base, Method=method, Error=None)
        except Exception as e:  # keep batch going on bad files
            kind, rec = (doc_type if doc_type != "auto" else "invoice"), {**base, "Error": str(e)}
        (chk if kind == "check" else inv).append(rec)
    return pd.DataFrame(inv), pd.DataFrame(chk)


def export_extraction(inv: pd.DataFrame, chk: pd.DataFrame, out_path: str):
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for name, d in (("Invoices", inv), ("Checks", chk)):
        if not d.empty:
            ws = wb.create_sheet(name)
            write_df(ws, d, money_cols=("Subtotal", "Tax", "Total", "Amount"))
    if not wb.sheetnames:
        wb.create_sheet("Empty")
    wb.save(out_path)


# ----------------------------------------------------------------------------
# 3. Reconciliation
# ----------------------------------------------------------------------------
def reconcile(internal: pd.DataFrame, bank: pd.DataFrame, date_tol: int = 3,
              amt_tol: float = 0.01, min_score: float = 0.35):
    """Greedy best-score matching. Candidates must be within amt_tol and date_tol
    days; ties are broken by description similarity and reference equality."""
    internal = internal.reset_index(drop=True).copy()
    bank = bank.reset_index(drop=True).copy()
    tol_c = int(round(amt_tol * 100))
    by_cents: dict[int, list[int]] = {}
    for j, a in bank.Amount.items():
        by_cents.setdefault(int(round(a * 100)), []).append(j)

    cands = []
    for i, r in internal.iterrows():
        c0 = int(round(r.Amount * 100))
        for c in range(c0 - tol_c, c0 + tol_c + 1):
            for j in by_cents.get(c, []):
                b = bank.loc[j]
                days = abs((r.Date - b.Date).days)
                if days > date_tol:
                    continue
                sim = difflib.SequenceMatcher(None, str(r.Description).lower(), str(b.Description).lower()).ratio()
                ref = 1.0 if r.Reference and r.Reference == b.Reference else 0.0
                score = 0.45 * sim + 0.30 * (1 - days / (date_tol + 1)) + 0.25 * ref + 0.0
                score += 0.20 if abs(r.Amount - b.Amount) < 0.005 else 0.0
                if score >= min_score or ref:
                    cands.append((score, i, j, days, sim))
    cands.sort(reverse=True)

    used_i, used_j, rows = set(), set(), []
    for score, i, j, days, sim in cands:
        if i in used_i or j in used_j:
            continue
        used_i.add(i), used_j.add(j)
        r, b = internal.loc[i], bank.loc[j]
        rows.append({
            "Internal Date": r.Date, "Internal Description": r.Description, "Internal Amount": r.Amount,
            "Bank Date": b.Date, "Bank Description": b.Description, "Bank Amount": b.Amount,
            "Day Diff": days, "Amount Diff": round(r.Amount - b.Amount, 2),
            "Similarity": round(sim, 2), "Confidence": round(min(score, 1.0), 2),
        })
    matched = pd.DataFrame(rows)
    un_int = internal.drop(index=list(used_i))[["Date", "Description", "Reference", "Amount"]]
    un_bank = bank.drop(index=list(used_j))[["Date", "Description", "Reference", "Amount"]]
    summary = pd.DataFrame([
        ("Internal transactions", len(internal)), ("Bank transactions", len(bank)),
        ("Matched pairs", len(matched)),
        ("Unmatched internal (not yet at bank)", len(un_int)),
        ("Unmatched bank (missing from books)", len(un_bank)),
        ("Internal total", round(internal.Amount.sum(), 2)), ("Bank total", round(bank.Amount.sum(), 2)),
        ("Unmatched internal total", round(un_int.Amount.sum(), 2)),
        ("Unmatched bank total", round(un_bank.Amount.sum(), 2)),
        ("Reconciled difference", round((internal.Amount.sum() - un_int.Amount.sum())
                                        - (bank.Amount.sum() - un_bank.Amount.sum()), 2)),
    ], columns=["Metric", "Value"])
    return matched, un_int.reset_index(drop=True), un_bank.reset_index(drop=True), summary


def export_reconciliation(matched, un_int, un_bank, summary, out_path: str):
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    write_df(ws, summary)
    for r in range(2, len(summary) + 2):
        ws.cell(row=r, column=1).font = Font(name=FONT)
    for name, d, money in (
        ("Matched", matched, ("Internal Amount", "Bank Amount", "Amount Diff")),
        ("Unmatched - Internal", un_int, ("Amount",)),
        ("Unmatched - Bank", un_bank, ("Amount",)),
    ):
        w = wb.create_sheet(name)
        if d.empty:
            w["A1"] = "None"
        else:
            write_df(w, d, money_cols=money, date_cols=("Date", "Internal Date", "Bank Date"))
    # Color-code summary difference
    last = len(summary) + 1
    ws.cell(row=last, column=2).font = Font(name=FONT, bold=True,
                                            color="008000" if summary.iloc[-1, 1] == 0 else "FF0000")
    wb.save(out_path)


# ----------------------------------------------------------------------------
# Sample data (so you can try the app immediately)
# ----------------------------------------------------------------------------
def make_sample_data(folder: str) -> tuple[str, str]:
    import random
    random.seed(7)
    Path(folder).mkdir(parents=True, exist_ok=True)
    cats_in = [("Product Sales", 1500, 6000), ("Service Revenue", 800, 3000), ("Interest Income", 20, 120)]
    cats_out = [("Salaries", 2500, 5000), ("Rent", 1800, 1800), ("Utilities", 90, 320),
                ("Software", 40, 400), ("Marketing", 150, 1200), ("Office Supplies", 20, 250)]
    rows, start = [], pd.Timestamp("2026-01-01")
    for d in range(0, 290, 1):
        day = start + pd.Timedelta(days=d)
        for _ in range(random.choice([0, 0, 1, 1, 2])):
            if random.random() < 0.4:
                c, lo, hi = random.choice(cats_in)
                amt = round(random.uniform(lo, hi), 2)
            else:
                c, lo, hi = random.choice(cats_out)
                amt = -round(random.uniform(lo, hi), 2)
            rows.append((day, f"{c} - ref {random.randint(1000, 9999)}", c, f"TX{len(rows) + 1:05d}", amt))
    ledger = pd.DataFrame(rows, columns=["Date", "Description", "Category", "Reference", "Amount"])
    # bank: shift dates 0-2 days, drop ~5%, add a few extras (fees)
    bank = ledger.copy()
    bank["Date"] = bank["Date"] + pd.to_timedelta([random.randint(0, 2) for _ in range(len(bank))], unit="D")
    bank["Description"] = bank["Description"].str.upper().str.replace(" - REF", " #", regex=False)
    bank["Reference"] = ""
    bank = bank.drop(bank.sample(frac=0.05, random_state=1).index)
    fees = pd.DataFrame([(start + pd.Timedelta(days=30 * k + 14), "MONTHLY SERVICE FEE", "Bank", "", -25.0)
                         for k in range(1, 9)], columns=bank.columns)
    bank = pd.concat([bank, fees]).sort_values("Date")[["Date", "Description", "Reference", "Amount"]]
    lp, bp = str(Path(folder) / "sample_ledger.csv"), str(Path(folder) / "sample_bank_statement.csv")
    ledger.to_csv(lp, index=False)
    bank.to_csv(bp, index=False)
    return lp, bp


# ----------------------------------------------------------------------------
# GUI
# ----------------------------------------------------------------------------
class App(tk.Tk if tk else object):
    def __init__(self):
        super().__init__()
        self.title("Financial Reporting & Document Processing Suite")
        self.geometry("1100x720")
        self.minsize(900, 600)
        style = ttk.Style(self)
        style.theme_use("clam")
        style.configure("TNotebook.Tab", padding=(16, 8), font=(FONT, 10, "bold"))
        style.configure("Accent.TButton", font=(FONT, 10, "bold"), padding=6)
        style.configure("Treeview.Heading", font=(FONT, 9, "bold"))
        self.vars = {k: tk.StringVar() for k in
                     ("ledger", "outdir", "company", "freq", "bank", "internal", "doctype", "dateTol", "amtTol")}
        self.vars["outdir"].set(str(Path.home() / "financial_reports"))
        self.vars["company"].set("My Company")
        self.vars["freq"].set("monthly")
        self.vars["doctype"].set("auto")
        self.vars["dateTol"].set("3")
        self.vars["amtTol"].set("0.01")
        self.doc_files: list[str] = []
        self.last_extract = (pd.DataFrame(), pd.DataFrame())
        self.last_recon = None

        self.status = tk.StringVar(value="Ready")
        nb = ttk.Notebook(self)
        nb.pack(fill="both", expand=True, padx=8, pady=(8, 0))
        for builder, title in ((self._tab_reports, "  Reports  "), (self._tab_extract, "  Data Extraction  "),
                               (self._tab_recon, "  Reconciliation  ")):
            f = ttk.Frame(nb, padding=10)
            nb.add(f, text=title)
            builder(f)
        bar = ttk.Frame(self)
        bar.pack(fill="x", padx=8, pady=6)
        ttk.Label(bar, textvariable=self.status).pack(side="left")
        self.progress = ttk.Progressbar(bar, mode="indeterminate", length=140)
        self.progress.pack(side="right")
        ttk.Button(bar, text="Create sample data", command=self.sample_data).pack(side="right", padx=10)

    # ----- widgets helpers
    def _file_row(self, parent, label, var, row, kind="open", types=(("CSV/Excel", "*.csv *.xlsx *.xls"),)):
        ttk.Label(parent, text=label).grid(row=row, column=0, sticky="w", pady=3)
        ttk.Entry(parent, textvariable=var, width=70).grid(row=row, column=1, sticky="ew", padx=6)

        def pick():
            p = filedialog.askdirectory() if kind == "dir" else filedialog.askopenfilename(filetypes=types)
            if p:
                var.set(p)
        ttk.Button(parent, text="Browse…", command=pick).grid(row=row, column=2)

    def _table(self, parent):
        wrap = ttk.Frame(parent)
        tv = ttk.Treeview(wrap, show="headings")
        sy, sx = ttk.Scrollbar(wrap, orient="vertical", command=tv.yview), ttk.Scrollbar(wrap, orient="horizontal", command=tv.xview)
        tv.configure(yscrollcommand=sy.set, xscrollcommand=sx.set)
        tv.grid(row=0, column=0, sticky="nsew")
        sy.grid(row=0, column=1, sticky="ns")
        sx.grid(row=1, column=0, sticky="ew")
        wrap.rowconfigure(0, weight=1)
        wrap.columnconfigure(0, weight=1)
        return wrap, tv

    @staticmethod
    def _fill(tv, df: pd.DataFrame, limit=2000):
        tv.delete(*tv.get_children())
        tv["columns"] = list(df.columns)
        for c in df.columns:
            tv.heading(c, text=c)
            tv.column(c, width=max(80, min(220, len(str(c)) * 11)), anchor="w")
        for row in df.head(limit).itertuples(index=False):
            vals = [("" if pd.isna(v) else (v.strftime("%Y-%m-%d") if isinstance(v, pd.Timestamp)
                     else (f"{v:,.2f}" if isinstance(v, float) else v))) for v in row]
            tv.insert("", "end", values=vals)

    def run_bg(self, func, done):
        self.progress.start(12)
        self.status.set("Working…")

        def worker():
            try:
                res = func()
                self.after(0, self._finish, done, res)
            except Exception as e:
                self.after(0, self._fail, e)
        threading.Thread(target=worker, daemon=True).start()

    def _finish(self, done, res):
        self.progress.stop()
        self.status.set("Done")
        done(res)

    def _fail(self, err):
        self.progress.stop()
        self.status.set("Error")
        messagebox.showerror("Error", str(err))

    # ----- tab 1: reports
    def _tab_reports(self, f):
        f.columnconfigure(1, weight=1)
        self._file_row(f, "Transaction ledger:", self.vars["ledger"], 0)
        self._file_row(f, "Output folder:", self.vars["outdir"], 1, kind="dir")
        ttk.Label(f, text="Company name:").grid(row=2, column=0, sticky="w", pady=3)
        ttk.Entry(f, textvariable=self.vars["company"], width=40).grid(row=2, column=1, sticky="w", padx=6)
        ttk.Label(f, text="Statement frequency:").grid(row=3, column=0, sticky="w", pady=3)
        ttk.Combobox(f, textvariable=self.vars["freq"], state="readonly", width=15,
                     values=["weekly", "monthly", "quarterly", "all"]).grid(row=3, column=1, sticky="w", padx=6)
        ttk.Button(f, text="Generate financial statements", style="Accent.TButton",
                   command=self.do_reports).grid(row=4, column=0, columnspan=2, sticky="w", pady=10)
        ttk.Label(f, text="Preview (last generated):").grid(row=5, column=0, sticky="w")
        wrap, self.rep_tv = self._table(f)
        wrap.grid(row=6, column=0, columnspan=3, sticky="nsew")
        f.rowconfigure(6, weight=1)

    def do_reports(self):
        if not self.vars["ledger"].get():
            return messagebox.showwarning("Missing input", "Choose a transaction ledger first.")

        def job():
            return generate_reports(self.vars["ledger"].get(), self.vars["freq"].get(),
                                    self.vars["outdir"].get(), self.vars["company"].get())

        def done(res):
            frames = []
            for f, (path, summ) in res.items():
                frames.append(summ.assign(Report=f.capitalize()))
            prev = pd.concat(frames)[["Report", "Period", "Revenue", "Expenses", "Net Income"]]
            self._fill(self.rep_tv, prev)
            messagebox.showinfo("Reports created", "\n".join(p for p, _ in res.values()))
        self.run_bg(job, done)

    # ----- tab 2: extraction
    def _tab_extract(self, f):
        f.columnconfigure(1, weight=1)
        top = ttk.Frame(f)
        top.grid(row=0, column=0, columnspan=3, sticky="ew")
        ttk.Button(top, text="Add files (PDF / JPG / PNG)…", command=self.add_docs).pack(side="left")
        ttk.Button(top, text="Clear", command=self.clear_docs).pack(side="left", padx=6)
        ttk.Label(top, text="Document type:").pack(side="left", padx=(20, 4))
        ttk.Combobox(top, textvariable=self.vars["doctype"], state="readonly", width=10,
                     values=["auto", "invoice", "check"]).pack(side="left")
        ttk.Button(top, text="Extract data", style="Accent.TButton", command=self.do_extract).pack(side="left", padx=14)
        ttk.Button(top, text="Export to Excel…", command=self.export_extract).pack(side="left")
        self.doc_label = ttk.Label(f, text="No files selected. OCR for scans/images needs Tesseract installed.")
        self.doc_label.grid(row=1, column=0, columnspan=3, sticky="w", pady=6)
        nb = ttk.Notebook(f)
        nb.grid(row=2, column=0, columnspan=3, sticky="nsew")
        f.rowconfigure(2, weight=1)
        self.inv_wrap, self.inv_tv = self._table(nb)
        self.chk_wrap, self.chk_tv = self._table(nb)
        nb.add(self.inv_wrap, text="Invoices")
        nb.add(self.chk_wrap, text="Checks")

    def add_docs(self):
        files = filedialog.askopenfilenames(filetypes=[("Documents", "*.pdf *.png *.jpg *.jpeg *.tif *.tiff *.bmp")])
        self.doc_files.extend(files)
        self.doc_label.config(text=f"{len(self.doc_files)} file(s) queued")

    def clear_docs(self):
        self.doc_files = []
        self.doc_label.config(text="No files selected.")

    def do_extract(self):
        if not self.doc_files:
            return messagebox.showwarning("Missing input", "Add at least one document.")
        self.run_bg(lambda: process_documents(self.doc_files, self.vars["doctype"].get()), self._show_extract)

    def _show_extract(self, res):
        self.last_extract = res
        inv, chk = res
        self._fill(self.inv_tv, inv)
        self._fill(self.chk_tv, chk)
        errs = sum(int(d["Error"].notna().sum()) for d in res if not d.empty and "Error" in d)
        self.status.set(f"Extracted {len(inv)} invoice(s), {len(chk)} check(s); {errs} error(s)")

    def export_extract(self):
        inv, chk = self.last_extract
        if inv.empty and chk.empty:
            return messagebox.showwarning("Nothing to export", "Run extraction first.")
        p = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
        if p:
            export_extraction(inv, chk, p)
            messagebox.showinfo("Saved", p)

    # ----- tab 3: reconciliation
    def _tab_recon(self, f):
        f.columnconfigure(1, weight=1)
        self._file_row(f, "Internal records:", self.vars["internal"], 0)
        self._file_row(f, "Bank statement:", self.vars["bank"], 1)
        opts = ttk.Frame(f)
        opts.grid(row=2, column=0, columnspan=3, sticky="w", pady=6)
        ttk.Label(opts, text="Date tolerance (days):").pack(side="left")
        ttk.Spinbox(opts, from_=0, to=30, textvariable=self.vars["dateTol"], width=4).pack(side="left", padx=(4, 16))
        ttk.Label(opts, text="Amount tolerance ($):").pack(side="left")
        ttk.Entry(opts, textvariable=self.vars["amtTol"], width=6).pack(side="left", padx=4)
        ttk.Button(opts, text="Reconcile", style="Accent.TButton", command=self.do_recon).pack(side="left", padx=16)
        ttk.Button(opts, text="Export to Excel…", command=self.export_recon).pack(side="left")
        self.sum_label = ttk.Label(f, text="", font=(FONT, 10, "bold"))
        self.sum_label.grid(row=3, column=0, columnspan=3, sticky="w")
        nb = ttk.Notebook(f)
        nb.grid(row=4, column=0, columnspan=3, sticky="nsew")
        f.rowconfigure(4, weight=1)
        self.rc = {}
        for key, title in (("matched", "Matched"), ("un_int", "Unmatched - Internal"),
                           ("un_bank", "Unmatched - Bank"), ("summary", "Summary")):
            wrap, tv = self._table(nb)
            nb.add(wrap, text=title)
            self.rc[key] = tv

    def do_recon(self):
        if not (self.vars["internal"].get() and self.vars["bank"].get()):
            return messagebox.showwarning("Missing input", "Select both files.")

        def job():
            a, b = load_transactions(self.vars["internal"].get()), load_transactions(self.vars["bank"].get())
            return reconcile(a, b, int(self.vars["dateTol"].get()), float(self.vars["amtTol"].get()))

        def done(res):
            self.last_recon = res
            for key, d in zip(("matched", "un_int", "un_bank", "summary"), res):
                self._fill(self.rc[key], d)
            m, ui, ub, s = res
            diff = s.iloc[-1, 1]
            self.sum_label.config(
                text=f"Matched {len(m)}  |  Unmatched internal {len(ui)}  |  Unmatched bank {len(ub)}  |  "
                     f"Reconciled difference: ${diff:,.2f}",
                foreground="#007a33" if diff == 0 else "#b00020")
        self.run_bg(job, done)

    def export_recon(self):
        if not self.last_recon:
            return messagebox.showwarning("Nothing to export", "Run reconciliation first.")
        p = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
        if p:
            export_reconciliation(*self.last_recon, p)
            messagebox.showinfo("Saved", p)

    def sample_data(self):
        folder = filedialog.askdirectory(title="Where should the sample files go?")
        if folder:
            lp, bp = make_sample_data(folder)
            self.vars["ledger"].set(lp)
            self.vars["internal"].set(lp)
            self.vars["bank"].set(bp)
            messagebox.showinfo("Sample data created", f"{lp}\n{bp}\n\nPaths were filled in for you.")


# ----------------------------------------------------------------------------
# CLI (for scheduled automation)
# ----------------------------------------------------------------------------
def cli(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Financial reporting suite (omit args to open the UI)")
    ap.add_argument("--report", metavar="LEDGER")
    ap.add_argument("--freq", default="all", choices=["weekly", "monthly", "quarterly", "all"])
    ap.add_argument("--company", default="My Company")
    ap.add_argument("--reconcile", nargs=2, metavar=("INTERNAL", "BANK"))
    ap.add_argument("--extract", nargs="+", metavar="FILE")
    ap.add_argument("--out", default="reports")
    ap.add_argument("--sample", metavar="FOLDER", help="write sample CSVs and exit")
    a = ap.parse_args(argv)
    if a.sample:
        print(*make_sample_data(a.sample), sep="\n")
        return 0
    if a.report:
        for f, (p, _) in generate_reports(a.report, a.freq, a.out, a.company).items():
            print(f"[{f}] {p}")
    if a.reconcile:
        res = reconcile(load_transactions(a.reconcile[0]), load_transactions(a.reconcile[1]))
        p = str(Path(a.out) / f"reconciliation_{datetime.now():%Y%m%d}.xlsx")
        export_reconciliation(*res, p)
        print(res[3].to_string(index=False))
        print("->", p)
    if a.extract:
        inv, chk = process_documents(a.extract)
        p = str(Path(a.out) / f"extracted_{datetime.now():%Y%m%d_%H%M}.xlsx")
        export_extraction(inv, chk, p)
        print(f"{len(inv)} invoices, {len(chk)} checks -> {p}")
    return 0


if __name__ == "__main__":
    if len(sys.argv) > 1:
        sys.exit(cli())
    if tk is None:
        sys.exit("Tkinter is not available. Install it (e.g. `sudo apt install python3-tk`) or use the CLI flags.")
    App().mainloop()
